#!/usr/bin/env python3
"""
Create a test Excel file for group import testing
"""

from openpyxl import Workbook

def create_test_excel():
    """Create a test Excel file with sample group data"""
    wb = Workbook()
    ws = wb.active
    
    # Headers (will be skipped during import)
    ws['A1'] = "Group Name"
    ws['B1'] = "Display ID"
    
    # Sample data - using different IDs to avoid conflicts
    test_data = [
        ("spotify 050", "050"),
        ("spotify 051", "051"),
        ("spotify 052", "052"),
        ("spotify 053", "053"),
        ("spotify 054", "054"),
    ]
    
    for i, (name, display_id) in enumerate(test_data, 2):
        ws[f'A{i}'] = name
        ws[f'B{i}'] = display_id
    
    # Save file
    wb.save("test_groups_import.xlsx")
    print("✅ Test Excel file created: test_groups_import.xlsx")
    
    # Show content
    print("\n📋 File content:")
    print("Row | Group Name    | Display ID")
    print("----|---------------|----------")
    print("1   | Group Name    | Display ID  (Header)")
    for i, (name, display_id) in enumerate(test_data, 2):
        print(f"{i}   | {name:<12} | {display_id}")

if __name__ == "__main__":
    create_test_excel()